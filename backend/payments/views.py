import io
import logging
from datetime import datetime, time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.contrib.auth import get_user_model
from django.db import transaction
from django.http import HttpResponse, Http404
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import generics, permissions
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.response import Response
from rest_framework.views import APIView

from groups.models import Group, GroupMembership

from .codes import generate_unique_receipt_code
from .models import TuitionCategory, TuitionTemplate, GroupTuition, StudentTuition, Payment, Enrollment, balances_for, balance_for
from .receipts import render_receipt_pdf
from .serializers import (
    TuitionCategorySerializer, TuitionTemplateSerializer,
    GroupTuitionSerializer, StudentTuitionSerializer, PaymentSerializer,
)

logger = logging.getLogger(__name__)


class IsAdmin(permissions.BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.role == 'admin'


def _student_label(user):
    return f'{user.first_name} {user.last_name}'.strip() or user.username


# ── Templates & categories ──────────────────────────────────────────────────

class TuitionCategoryListCreateView(generics.ListCreateAPIView):
    serializer_class = TuitionCategorySerializer
    permission_classes = (IsAdmin,)

    def get_queryset(self):
        return TuitionCategory.objects.filter(academy=self.request.user.academy)

    def perform_create(self, serializer):
        name = serializer.validated_data['name']
        if TuitionCategory.objects.filter(academy=self.request.user.academy, name__iexact=name).exists():
            raise ValidationError({'name': "Bu nomdagi kategoriya allaqachon mavjud."})
        serializer.save(academy=self.request.user.academy)


class TuitionTemplateListCreateView(generics.ListCreateAPIView):
    serializer_class = TuitionTemplateSerializer
    permission_classes = (IsAdmin,)

    def get_queryset(self):
        return TuitionTemplate.objects.filter(academy=self.request.user.academy).select_related('category')

    def perform_create(self, serializer):
        category = serializer.validated_data['category']
        name = serializer.validated_data['name']
        if category.academy_id != self.request.user.academy_id:
            raise PermissionDenied("Bu kategoriya boshqa akademiyaga tegishli.")
        if TuitionTemplate.objects.filter(academy=self.request.user.academy, category=category, name__iexact=name).exists():
            raise ValidationError({'name': "Bu nomdagi shablon shu kategoriyada allaqachon mavjud."})
        serializer.save(academy=self.request.user.academy)


class GroupTuitionView(APIView):
    """Assign/change which template a group uses."""

    permission_classes = (IsAdmin,)

    def put(self, request, pk):
        group = get_object_or_404(Group, pk=pk, teacher__academy=request.user.academy)
        template = get_object_or_404(TuitionTemplate, pk=request.data.get('template'), academy=request.user.academy)
        obj, _ = GroupTuition.objects.update_or_create(group=group, defaults={'template': template})
        return Response(GroupTuitionSerializer(obj).data)


class StudentTuitionOverrideView(APIView):
    """Get/set a per-student custom price for one group."""

    permission_classes = (IsAdmin,)

    def get(self, request, pk, group_id):
        override = StudentTuition.objects.filter(student_id=pk, group_id=group_id).first()
        if not override:
            return Response(None)
        return Response(StudentTuitionSerializer(override).data)

    def put(self, request, pk, group_id):
        group = get_object_or_404(Group, pk=group_id, teacher__academy=request.user.academy)
        get_object_or_404(GroupMembership, student_id=pk, group=group)
        custom_price = request.data.get('custom_price')
        if custom_price in (None, ''):
            StudentTuition.objects.filter(student_id=pk, group=group).delete()
            return Response(None)
        obj, _ = StudentTuition.objects.update_or_create(
            student_id=pk, group=group, defaults={'custom_price': custom_price},
        )
        return Response(StudentTuitionSerializer(obj).data)


# ── Balances & payments ──────────────────────────────────────────────────────

class BalancesListView(APIView):
    """Admin-only: every student's expected/paid/balance, grouped by Group so
    the admin can scan per-cohort instead of one long flat list — an academy
    with many groups/students otherwise produces an unreadable wall of rows."""

    permission_classes = (IsAdmin,)

    def get(self, request):
        rows = balances_for(request.user.academy)

        group_filter = request.query_params.get('group')
        if group_filter:
            rows = [r for r in rows if r['group'].id == int(group_filter)]
        only_debt = request.query_params.get('only_debt') == 'true'
        if only_debt:
            rows = [r for r in rows if r['balance'] > 0]

        by_group = {}
        for r in rows:
            entry = by_group.setdefault(r['group'].id, {
                'group_id': r['group'].id, 'group_name': r['group'].name,
                'rows': [], 'debtor_count': 0, 'total_debt': 0,
            })
            entry['rows'].append({
                'student_id': r['student'].id,
                'student_name': _student_label(r['student']),
                'price': r['price'],
                'expected': r['expected'],
                'paid': r['paid'],
                'paid_this_month': r['paid_this_month'],
                'balance': r['balance'],
                'joined_at': r['started_at'],
                'ended_at': r['ended_at'],
                'next_due_date': r['next_due_date'],
            })
            if r['balance'] > 0:
                entry['debtor_count'] += 1
                entry['total_debt'] += r['balance']

        groups = sorted(by_group.values(), key=lambda g: -g['total_debt'])
        for g in groups:
            g['rows'].sort(key=lambda r: -r['balance'])

        return Response({
            'groups': groups,
            'summary': {
                'debtor_count': sum(1 for r in rows if r['balance'] > 0),
                'total_debt': sum(r['balance'] for r in rows if r['balance'] > 0),
            },
        })


class BalancesExportExcelView(APIView):
    """Admin-only: the whole (or one group's) balance list as an .xlsx —
    for bookkeeping, printing, archiving a point-in-time snapshot, or
    sharing with someone who has no login. Mirrors the .xlsx pattern
    already used for attendance/exam exports in the groups app."""

    permission_classes = (IsAdmin,)

    def get(self, request):
        rows = balances_for(request.user.academy)
        group_filter = request.query_params.get('group')
        if group_filter:
            rows = [r for r in rows if r['group'].id == int(group_filter)]
        only_debt = request.query_params.get('only_debt') == 'true'
        if only_debt:
            rows = [r for r in rows if r['balance'] > 0]
        rows.sort(key=lambda r: (r['group'].name, -r['balance']))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Balans'

        academy_name = request.user.academy.name if request.user.academy else ''
        today = timezone.localdate().strftime('%d.%m.%Y')

        ws.merge_cells('A1:I1')
        ws['A1'] = f"{academy_name} — to'lovlar balansi ({today} holatiga)"
        ws['A1'].font = Font(bold=True, size=13)
        ws.row_dimensions[1].height = 24

        headers = ["№", 'Guruh', "O'quvchi", 'Qo\'shilgan sana', 'Oylik narx', 'Bu oy to\'lagan', 'Jami to\'lagan', 'Kerak edi', 'Balans (qarz)']
        widths  = [5, 24, 24, 16, 14, 14, 14, 14, 16]
        header_fill = PatternFill('solid', fgColor='1F618D')
        white_bold  = Font(bold=True, color='FFFFFF', size=10)
        center      = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for j, (label, width) in enumerate(zip(headers, widths)):
            col = get_column_letter(j + 1)
            cell = ws[f'{col}2']
            cell.value = label; cell.font = white_bold; cell.fill = header_fill; cell.alignment = center
            ws.column_dimensions[col].width = width
        ws.row_dimensions[2].height = 30

        debt_font = Font(color='991B1B', bold=True)
        ok_font   = Font(color='166534', bold=True)

        for i, r in enumerate(rows):
            row = i + 3
            joined = r['started_at']
            joined = joined.date() if hasattr(joined, 'date') else joined
            values = [
                i + 1, r['group'].name, _student_label(r['student']),
                joined.strftime('%d.%m.%Y'), r['price'], r['paid_this_month'],
                r['paid'], r['expected'], r['balance'],
            ]
            for j, value in enumerate(values):
                col = get_column_letter(j + 1)
                cell = ws[f'{col}{row}']
                cell.value = value
                if col == 'A':
                    cell.alignment = center
            balance_cell = ws[f'I{row}']
            balance_cell.font = debt_font if r['balance'] > 0 else ok_font

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        filename = 'tolovlar_balansi.xlsx' if not group_filter else f"{rows[0]['group'].name.replace(' ', '_')}_balans.xlsx" if rows else 'balans.xlsx'
        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class RecordPaymentView(APIView):
    permission_classes = (IsAdmin,)

    def post(self, request):
        student_id = request.data.get('student')
        group_id   = request.data.get('group')
        amount     = request.data.get('amount')
        method     = request.data.get('method')

        if method not in Payment.Method.values:
            return Response({'detail': "To'lov usuli noto'g'ri."}, status=400)
        try:
            amount = int(amount)
        except (TypeError, ValueError):
            return Response({'detail': "Summa noto'g'ri."}, status=400)
        if amount < 1:
            return Response({'detail': "Summa 1 so'mdan katta bo'lishi kerak."}, status=400)

        paid_at = timezone.now()
        paid_at_raw = request.data.get('paid_at')
        if paid_at_raw:
            try:
                d = datetime.strptime(str(paid_at_raw)[:10], '%Y-%m-%d').date()
            except ValueError:
                return Response({'detail': "Sana noto'g'ri."}, status=400)
            if d > timezone.localdate():
                return Response({'detail': "Sana kelajakda bo'lishi mumkin emas."}, status=400)
            # Keep the current time-of-day for a same-day entry (matches
            # "now"); backdated entries just get midnight — the exact time
            # of a past payment usually isn't known anyway.
            paid_at = timezone.make_aware(datetime.combine(d, time.min)) if d != timezone.localdate() else paid_at

        group = get_object_or_404(Group, pk=group_id, teacher__academy=request.user.academy)
        # Enrollment (not GroupMembership) so a payment can still be recorded
        # against a student who has since left the group — settling old debt
        # shouldn't require re-adding them.
        if not Enrollment.objects.filter(student_id=student_id, group=group).exists():
            raise Http404("Bu o'quvchi bu guruhga hech qachon a'zo bo'lmagan.")
        student = get_object_or_404(get_user_model(), pk=student_id)

        with transaction.atomic():
            payment = Payment.objects.create(
                student=student, group=group, amount=amount, method=method,
                note=request.data.get('note', ''), receipt_code=generate_unique_receipt_code(),
                recorded_by=request.user, paid_at=paid_at,
            )

        try:
            from .notifications import send_payment_receipt
            send_payment_receipt(payment)
        except Exception as e:
            logger.error('Failed to send payment receipt over Telegram: %s', e)

        return Response(PaymentSerializer(payment).data, status=201)


class PaymentHistoryView(APIView):
    permission_classes = (IsAdmin,)

    def get(self, request):
        qs = Payment.objects.filter(group__teacher__academy=request.user.academy).select_related('student', 'group')
        student_filter = request.query_params.get('student')
        if student_filter:
            qs = qs.filter(student_id=student_filter)
        group_filter = request.query_params.get('group')
        if group_filter:
            qs = qs.filter(group_id=group_filter)

        page      = max(1, int(request.query_params.get('page', 1)))
        page_size = max(1, min(50, int(request.query_params.get('page_size', 20))))
        total     = qs.count()
        pages     = max(1, (total + page_size - 1) // page_size)
        page      = min(page, pages)
        payments  = list(qs[(page - 1) * page_size: page * page_size])

        return Response({
            'results': PaymentSerializer(payments, many=True).data,
            'total': total, 'pages': pages, 'page': page,
        })


class PaymentVoidView(APIView):
    """Admin-only: cancel a wrongly-entered payment. A hard delete (not a
    reversing ledger entry like coins/purchases use) — a Payment here is
    already a manually-entered record of a real-world event, so an admin
    correcting a typo is deleting a mistake, not reversing a real one."""

    permission_classes = (IsAdmin,)

    def delete(self, request, pk):
        payment = get_object_or_404(Payment, pk=pk, group__teacher__academy=request.user.academy)
        payment.delete()
        return Response(status=204)


class PaymentReceiptView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request, pk):
        payment = get_object_or_404(Payment.objects.select_related('student', 'group', 'group__teacher__academy'), pk=pk)
        is_owner = payment.student_id == request.user.id
        is_admin_here = request.user.role == 'admin' and payment.group.teacher.academy_id == request.user.academy_id
        is_parent = False
        if request.user.role == 'parent':
            from users.models import ParentStudent
            is_parent = ParentStudent.objects.filter(parent=request.user, student_id=payment.student_id).exists()
        if not (is_owner or is_admin_here or is_parent):
            return Response({'detail': "Ruxsat yo'q."}, status=403)

        pdf_bytes = render_receipt_pdf(payment)
        response = HttpResponse(pdf_bytes, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="receipt-{payment.receipt_code}.pdf"'
        return response


class MyPaymentsView(APIView):
    """Student/parent self-service: own (or child's) balance per group + history."""

    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request):
        if request.user.role == 'student':
            students = [request.user]
        elif request.user.role == 'parent':
            from users.models import ParentStudent
            students = [ps.student for ps in ParentStudent.objects.filter(parent=request.user).select_related('student')]
        else:
            return Response({'detail': "Faqat o'quvchi yoki ota-ona ko'ra oladi."}, status=403)

        enrollments = (
            Enrollment.objects.filter(student__in=students)
            .select_related('student', 'group', 'group__tuition__template')
            .order_by('student_id', 'group_id', '-started_at')
        )
        latest = {}
        for e in enrollments:
            latest.setdefault((e.student_id, e.group_id), e)

        results = []
        for e in latest.values():
            results.append({
                'student_id': e.student_id,
                'student_name': _student_label(e.student),
                'group_id': e.group.id,
                'group_name': e.group.name,
                'balance': balance_for(e.student, e.group, e),
            })

        payments = Payment.objects.filter(student__in=students).select_related('student', 'group')[:20]
        return Response({
            'balances': results,
            'recent_payments': PaymentSerializer(payments, many=True).data,
        })
